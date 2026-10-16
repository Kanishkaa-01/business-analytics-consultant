import io
import pytest
import openpyxl
from fastapi.testclient import TestClient
from app.main import app
from app.core.database import get_db, SessionLocal
from app.models.db_models import AnalysisRun, Dataset, DecisionRecord, SimulationScenario
from app.services.report_generator import ReportGeneratorService


@pytest.fixture
def mock_report_data():
    run = AnalysisRun(
        id="test-run-rep-001",
        dataset_id="test-ds-rep-001",
        target_column="Revenue",
        date_column="Date",
        status="completed",
        execution_time_ms=125.4,
        data_quality={
            "overall_quality_score": 94.5,
            "reliability_assessment": "High Reliability",
            "completeness_pct": 98.2,
            "validity_pct": 99.0,
            "column_metrics": [
                {
                    "column_name": "Revenue",
                    "inferred_type": "Numeric",
                    "completeness_pct": 100.0,
                    "validity_pct": 100.0,
                    "null_count": 0,
                    "unique_count": 150,
                    "potential_issues": []
                }
            ]
        },
        business_health={
            "overall_health_score": 82,
            "overall_status": "Stable",
            "executive_summary": "Financial performance is stable with positive revenue growth.",
            "kpis": [
                {
                    "name": "Revenue",
                    "current_value": 450000.0,
                    "target_or_baseline": 420000.0,
                    "variance_pct": 7.1,
                    "unit": "$",
                    "status": "healthy",
                    "interpretation": "Exceeding target"
                }
            ]
        },
        diagnostics={
            "focal_metric": "Revenue",
            "total_variance": 30000.0,
            "variance_pct": 7.1,
            "diagnostic_narrative": "Observed positive growth across major product lines.",
            "findings": [
                {
                    "title": "Revenue Expansion",
                    "type": "Growth",
                    "metric": "Revenue",
                    "observed_value": 450000.0,
                    "comparison": "Baseline",
                    "difference": 30000.0,
                    "possible_reason": "Strong channel demand",
                    "recommended_action": "Expand product supply"
                }
            ],
            "variance_breakdown": [
                {
                    "dimension": "Category",
                    "segment": "Electronics",
                    "baseline_value": 200000.0,
                    "actual_value": 230000.0,
                    "variance": 30000.0,
                    "variance_pct": 15.0,
                    "contribution_share_pct": 80.0,
                    "direction": "Lift"
                }
            ]
        },
        root_cause={
            "primary_root_cause": "Electronics demand increase",
            "statistical_confidence_level": 88.0,
            "factor_contributions": [
                {
                    "factor_name": "Product Demand",
                    "factor_type": "Commercial",
                    "attribution_percentage": 75.0,
                    "direction": "Positive Impact",
                    "evidence_statement": "Correlation +0.81"
                }
            ],
            "pareto_breakdown": [
                {
                    "entity": "SKU-1",
                    "loss_or_impact": 20000.0,
                    "cumulative_percentage": 66.7,
                    "is_vital_few": True
                }
            ]
        },
        risks_opportunities={
            "total_exposure_at_risk": 15000.0,
            "total_opportunity_upside": 40000.0,
            "risk_profile_rating": "Moderate",
            "risks": [],
            "opportunities": []
        },
        evidence_chains=[
            {
                "id": "EV-01",
                "title": "Electronics Channel Expansion",
                "target_metric": "Revenue",
                "observed_phenomenon": "Electronics growth +$30k",
                "diagnostic_pathway": ["Category", "Electronics"],
                "underlying_mechanism": "Direct e-commerce uptick",
                "statistical_proof": "p-value < 0.01",
                "associated_recommendation_id": "REC-001",
                "nodes": []
            }
        ],
        recommendations=[
            {
                "id": "REC-001",
                "title": "Scale Supply for Electronics",
                "urgency": "Immediate (0–30 days)",
                "category": "Growth Opportunity",
                "mechanism_of_action": "Increase inventory buffer by 15%",
                "rationale": "High sales velocity",
                "prerequisites": ["Check supplier lead times"],
                "expected_financial_impact": 25000.0,
                "implementation_effort": "Low",
                "confidence_score_pct": 85.0,
                "impact_score": 8.0,
                "confidence_score": 8.5,
                "ease_score": 8.0,
                "priority_score": 82.5,
                "priority_tier": "High Priority",
                "priority_level": "High Priority",
                "tradeoffs_and_risks": "Carrying cost increase",
                "status": "Pending Review",
                "what_to_do": "Increase inventory buffer by 15%",
                "why": "High sales velocity",
                "supporting_finding": "Electronics revenue grew +$30k",
                "decision_status": "Pending Review"
            }
        ]
    )

    dataset = Dataset(
        id="test-ds-rep-001",
        name="Q3 Performance",
        filename="q3_perf.csv",
        file_path="",
        file_size_bytes=1024,
        file_hash="abc123hash",
        row_count=500,
        column_count=8
    )

    decisions = [
        DecisionRecord(
            id="dec-001",
            analysis_id="test-run-rep-001",
            recommendation_id="REC-001",
            action_title="Scale Supply for Electronics",
            decision_status="Accepted",
            stakeholder="COO",
            notes_rationale="Approved budget for inventory expansion.",
            target_date="2026-11-15",
            expected_impact_dollars=25000.0
        )
    ]

    scenarios = [
        SimulationScenario(
            id="scen-001",
            analysis_id="test-run-rep-001",
            name="Aggressive Expansion",
            parameters={"volume_change_pct": 10.0},
            projected_outcomes={"simulated_metric_value": 495000.0}
        )
    ]

    return run, dataset, decisions, scenarios


def test_pdf_report_generation(mock_report_data):
    run, dataset, decisions, scenarios = mock_report_data
    pdf_bytes = ReportGeneratorService.generate_pdf_report(run, dataset, decisions, scenarios)
    
    assert isinstance(pdf_bytes, bytes)
    assert len(pdf_bytes) > 1000
    # PDF signature
    assert pdf_bytes.startswith(b"%PDF")


def test_excel_report_generation(mock_report_data):
    run, dataset, decisions, scenarios = mock_report_data
    excel_bytes = ReportGeneratorService.generate_excel_report(run, dataset, decisions, scenarios)
    
    assert isinstance(excel_bytes, bytes)
    assert len(excel_bytes) > 1000
    # Zip signature for .xlsx
    assert excel_bytes.startswith(b"PK\x03\x04")

    # Load with openpyxl to inspect worksheets
    wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
    sheet_names = wb.sheetnames

    expected_sheets = [
        "Summary",
        "Data Quality",
        "Business Metrics",
        "Findings",
        "Trends",
        "Anomalies",
        "Segment Comparisons",
        "Recommendations",
        "Scenario Comparison",
        "Decision Notes"
    ]
    for s in expected_sheets:
        assert s in sheet_names

    # Check Summary sheet content
    ws_sum = wb["Summary"]
    assert ws_sum.cell(row=1, column=1).value == "Business Decision & Analytics Summary"

    # Check Recommendations sheet content
    ws_rec = wb["Recommendations"]
    assert ws_rec.cell(row=4, column=3).value == "Scale Supply for Electronics"
    assert ws_rec.cell(row=4, column=4).value == "High Priority"


def test_api_export_endpoints(mock_report_data):
    run, dataset, decisions, scenarios = mock_report_data
    db = SessionLocal()
    try:
        # Save mock objects to db
        db.add(dataset)
        db.add(run)
        for d in decisions:
            db.add(d)
        for s in scenarios:
            db.add(s)
        db.commit()

        client = TestClient(app)

        # Test PDF export route
        pdf_res = client.get(f"/api/export/{run.id}/pdf-report")
        assert pdf_res.status_code == 200
        assert pdf_res.headers["content-type"] == "application/pdf"
        assert pdf_res.content.startswith(b"%PDF")

        # Test Excel export route
        excel_res = client.get(f"/api/export/{run.id}/excel-report")
        assert excel_res.status_code == 200
        assert "spreadsheetml" in excel_res.headers["content-type"]
        assert excel_res.content.startswith(b"PK\x03\x04")

        # Test 404 for unknown run
        res_404 = client.get("/api/export/non-existent-run-id/pdf-report")
        assert res_404.status_code == 404

    finally:
        # Cleanup
        db.query(SimulationScenario).filter(SimulationScenario.analysis_id == run.id).delete()
        db.query(DecisionRecord).filter(DecisionRecord.analysis_id == run.id).delete()
        db.query(AnalysisRun).filter(AnalysisRun.id == run.id).delete()
        db.query(Dataset).filter(Dataset.id == dataset.id).delete()
        db.commit()
        db.close()


def test_pdf_report_data_quality_mapping_and_no_zero_range():
    """Verify that Data Quality, Trends, and Anomalies map real pipeline fields without blank cells or [0.0, 0.0]."""
    import pymupdf

    run = AnalysisRun(
        id="test-run-pdf-002",
        dataset_id="test-ds-pdf-002",
        target_column="Revenue",
        date_column="Date",
        status="completed",
        execution_time_ms=95.0,
        data_quality={
            "overall_quality_score": 98.0,
            "reliability_assessment": "High Reliability",
            "completeness_ratio": 1.0,
            "validity_ratio": 1.0,
            "column_metrics": [
                {
                    "column": "Units_Sold",
                    "completeness_score": 1.0,
                    "validity_score": 0.985,
                    "outlier_count": 2,
                    "outlier_percentage": 1.5,
                }
            ],
            "issues": [
                {
                    "column": "Units_Sold",
                    "dimension": "Distribution",
                    "severity": "Warning",
                    "description": "Minor outlier concentration."
                }
            ]
        },
        business_health={
            "overall_health_score": 85,
            "overall_status": "Healthy",
            "executive_summary": "Solid quarter with growth in core channels.",
            "kpis": [
                {
                    "name": "Revenue",
                    "current_value": 500000.0,
                    "target_or_baseline": 480000.0,
                    "variance_pct": 4.17,
                    "unit": "$",
                    "status": "healthy"
                }
            ]
        },
        diagnostics={
            "focal_metric": "Revenue",
            "total_variance": 20000.0,
            "variance_pct": 4.17,
            "diagnostic_narrative": "Observed positive growth in wholesale accounts.",
            "findings": [
                {
                    "title": "Strong Wholesale Lift",
                    "metric": "Revenue",
                    "observed_value": 500000.0,
                    "comparison": "Baseline",
                    "difference": 20000.0
                }
            ],
            "variance_breakdown": [
                {
                    "dimension": "Channel",
                    "segment": "Wholesale",
                    "baseline_value": 200000.0,
                    "actual_value": 220000.0,
                    "variance": 20000.0,
                    "variance_pct": 10.0,
                    "contribution_share_pct": 100.0,
                    "direction": "Lift"
                }
            ]
        },
        root_cause={
            "primary_root_cause": "Wholesale contract renewal",
            "statistical_confidence_level": 92.0,
            "factor_contributions": [
                {
                    "factor_name": "Contract Renewals",
                    "factor_type": "Commercial",
                    "attribution_percentage": 85.0,
                    "direction": "Positive",
                    "evidence_statement": "Wholesale volume +10%"
                }
            ]
        },
        risks_opportunities={},
        evidence_chains=[],
        recommendations=[]
    )

    dataset = Dataset(
        id="test-ds-pdf-002",
        name="Sales Ledger Q3",
        filename="sales_ledger.csv",
        file_path="",
        file_size_bytes=2048,
        file_hash="hash456def",
        row_count=1000,
        column_count=5,
        schema_metadata={
            "columns": [
                {
                    "name": "Units_Sold",
                    "inferred_type": "numeric"
                }
            ]
        }
    )

    pdf_bytes = ReportGeneratorService.generate_pdf_report(run, dataset, [], [])
    assert isinstance(pdf_bytes, bytes)
    assert pdf_bytes.startswith(b"%PDF")

    # Extract text with pymupdf and verify accurate data rendering
    doc = pymupdf.open(stream=pdf_bytes, filetype="pdf")
    full_text = "\n".join(page.get_text() for page in doc)

    # 1. Column name and type are populated
    assert "Units_Sold" in full_text
    assert "Numeric" in full_text
    # 2. Completeness and validity are properly formatted (100.0%, 98.5%)
    assert "100.0%" in full_text
    assert "98.5%" in full_text
    # 3. Quality issue is mapped
    assert "Distribution (Warning)" in full_text
    # 4. Expected range does not display misleading [0.0, 0.0]
    assert "[0.0, 0.0]" not in full_text
    # 5. Segment breakdown columns and values are populated
    assert "Wholesale" in full_text
    assert "Positive Gain" in full_text


def test_stored_analyses_backend_pagination():
    """Verify stored analyses endpoint pagination metadata, sorting, and backward compatibility."""
    client = TestClient(app)
    db = SessionLocal()

    created_runs = []
    created_datasets = []
    try:
        for i in range(3):
            ds = Dataset(
                id=f"test-ds-pg-{i}",
                name=f"Dataset {chr(65+i)}",
                filename=f"ds_{i}.csv",
                file_path="",
                file_size_bytes=100,
                file_hash=f"hash_{i}",
                row_count=50,
                column_count=3
            )
            run = AnalysisRun(
                id=f"test-run-pg-{i}",
                dataset_id=ds.id,
                target_column="Sales",
                analysis_name=f"Analysis Run {chr(65+i)}",
                status="completed",
                execution_time_ms=50.0 + i * 10,
                business_health={"overall_health_score": 70.0 + i * 5}
            )
            db.add(ds)
            db.add(run)
            created_datasets.append(ds)
            created_runs.append(run)
        db.commit()

        # 1. Paginated request
        res_p1 = client.get("/api/analyses?page=1&page_size=2&sort_by=name")
        assert res_p1.status_code == 200
        data_p1 = res_p1.json()
        assert "items" in data_p1
        assert "total" in data_p1
        assert "page" in data_p1
        assert "page_size" in data_p1
        assert "total_pages" in data_p1

        assert data_p1["page"] == 1
        assert data_p1["page_size"] == 2
        assert data_p1["total"] >= 3
        assert len(data_p1["items"]) == 2
        assert data_p1["total_pages"] >= 2

        # 2. Page 2 request
        res_p2 = client.get("/api/analyses?page=2&page_size=2&sort_by=name")
        assert res_p2.status_code == 200
        data_p2 = res_p2.json()
        assert data_p2["page"] == 2
        assert len(data_p2["items"]) >= 1

        # 3. Backward compatibility (unpaginated request returns raw list)
        res_unp = client.get("/api/analyses")
        assert res_unp.status_code == 200
        data_unp = res_unp.json()
        assert isinstance(data_unp, list)
        assert len(data_unp) >= 3

        # 4. Alias route works identically
        res_alias = client.get("/api/analysis/stored?page=1&page_size=2")
        assert res_alias.status_code == 200
        assert "items" in res_alias.json()

    finally:
        for r in created_runs:
            db.query(AnalysisRun).filter(AnalysisRun.id == r.id).delete()
        for d in created_datasets:
            db.query(Dataset).filter(Dataset.id == d.id).delete()
        db.commit()
        db.close()

