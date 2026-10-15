from pathlib import Path
import pymupdf
import docx
from docx.shared import Inches, Pt, RGBColor


def generate_document_samples():
    sample_dir = Path(__file__).parent
    sample_dir.mkdir(parents=True, exist_ok=True)

    # -------------------------------------------------------------
    # 1. Generate retail_q3_executive_memo.pdf
    # -------------------------------------------------------------
    pdf_path = sample_dir / "retail_q3_executive_memo.pdf"
    doc = pymupdf.open()

    # Page 1: Executive Overview & Commercial Policy
    page1 = doc.new_page(width=595, height=842)  # A4
    page1.insert_text((50, 60), "Q3 Retail Operations & Performance Review", fontsize=18, fontname="helv", color=(0.1, 0.2, 0.4))
    page1.insert_text((50, 85), "Executive Management Memorandum — Author: Commercial Steering Group", fontsize=10, fontname="helv", color=(0.4, 0.4, 0.5))

    page1.insert_text((50, 120), "1. Executive Performance Summary", fontsize=14, fontname="helv", color=(0.15, 0.25, 0.45))
    text_p1 = (
        "During the third quarter of 2025, overall gross sales revenue reached $52.4M across all direct and "
        "marketplace channels. While customer acquisition in the consumer retail segment demonstrated resilience, "
        "overall net margin capture deteriorated by 14.2% compared to the prior half baseline.\n\n"
        "The primary operational drag stems from marketplace discounting in Consumer Electronics, where uncoordinated "
        "promotional markdown rates surged to 31.4% in the second half of the quarter. Furthermore, European distribution "
        "centers reported average delivery lead times escalating to 8.5 days due to regional customs bottlenecks."
    )
    rect1 = pymupdf.Rect(50, 140, 545, 260)
    page1.insert_textbox(rect1, text_p1, fontsize=10, fontname="helv", color=(0.2, 0.2, 0.2))

    page1.insert_text((50, 280), "2. Commercial Channel Performance Table", fontsize=14, fontname="helv", color=(0.15, 0.25, 0.45))

    # Draw a clean table on Page 1
    table_y = 310
    headers = ["Sales Channel", "Gross Sales", "Avg Discount", "Net Margin", "Status"]
    col_widths = [110, 95, 95, 95, 90]
    
    # Header row
    x = 50
    for h, w in zip(headers, col_widths):
        page1.draw_rect(pymupdf.Rect(x, table_y, x + w, table_y + 22), color=(0.8, 0.8, 0.8), fill=(0.93, 0.95, 0.98))
        page1.insert_text((x + 8, table_y + 15), h, fontsize=9, fontname="helv", color=(0.1, 0.2, 0.4))
        x += w

    rows = [
        ["Direct Online", "$24.8M", "8.5%", "$4.2M", "Healthy"],
        ["Marketplace", "$18.6M", "31.4%", "$1.1M", "Vulnerable"],
        ["Wholesale B2B", "$9.0M", "5.0%", "$1.8M", "Stable"],
    ]

    for r_idx, row in enumerate(rows):
        y = table_y + 22 + (r_idx * 20)
        x = 50
        for val, w in zip(row, col_widths):
            page1.draw_rect(pymupdf.Rect(x, y, x + w, y + 20), color=(0.85, 0.85, 0.85), fill=(1, 1, 1))
            page1.insert_text((x + 8, y + 14), val, fontsize=9, fontname="helv", color=(0.2, 0.2, 0.2))
            x += w

    # Page 2: Recommendations and Risk Exposure
    page2 = doc.new_page(width=595, height=842)
    page2.insert_text((50, 60), "3. Operational Remediation & Governance Directives", fontsize=14, fontname="helv", color=(0.15, 0.25, 0.45))
    text_p2 = (
        "To protect portfolio unit economics, management mandates the immediate rollout of an algorithmic "
        "discount floor across all third-party marketplace channels, capping discretionary sales rep concessions at 15.0%.\n\n"
        "Without this intervention, commercial downside exposure is estimated at $480,000 per month in lost gross margin. "
        "Concurrently, the European supply chain steering committee will audit warehouse fulfillment workflows to restore the "
        "standard 48-hour delivery SLA."
    )
    rect2 = pymupdf.Rect(50, 85, 545, 220)
    page2.insert_textbox(rect2, text_p2, fontsize=10, fontname="helv", color=(0.2, 0.2, 0.2))

    doc.save(str(pdf_path))
    doc.close()
    print(f"Generated {pdf_path.name}")

    # -------------------------------------------------------------
    # 2. Generate q3_financial_audit_review.docx
    # -------------------------------------------------------------
    docx_path = sample_dir / "q3_financial_audit_review.docx"
    doc_docx = docx.Document()

    doc_docx.core_properties.title = "Q3 Internal Audit & Financial Control Review"
    doc_docx.core_properties.author = "Corporate Financial Audit Group"

    # Title
    p_title = doc_docx.add_heading("Q3 Financial Audit & Revenue Reconciliation", level=0)

    # Section 1
    doc_docx.add_heading("1. Audit Summary & Accounting Discrepancies", level=1)
    doc_docx.add_paragraph(
        "The Corporate Internal Audit Group has completed its periodic examination of third-quarter commercial revenue. "
        "Following adjustments for promotional allowances, sales returns, and deferred marketplace commissions, "
        "recognized Net Commercial Revenue stands at $49.8M."
    )
    doc_docx.add_paragraph(
        "Notice of Data Variance: Operating sales memoranda previously indicated gross booking volumes of $52.4M. "
        "The $2.6M variance represents unaccrued marketplace markdown charges and regional freight equalization deductions."
    )

    # Section 2: Table
    doc_docx.add_heading("2. Regional Revenue Reconciliation Table", level=1)
    table = doc_docx.add_table(rows=1, cols=4)
    table.style = "Table Grid"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Operating Region"
    hdr_cells[1].text = "Reported Revenue"
    hdr_cells[2].text = "Adjusted Margin"
    hdr_cells[3].text = "Audit Status"

    data = [
        ["North America", "$21.5M", "$3.8M", "Verified"],
        ["Europe", "$14.2M", "$1.6M", "Audit Adjustment"],
        ["Asia Pacific", "$9.6M", "$1.9M", "Verified"],
        ["Latin America", "$4.5M", "$0.8M", "Verified"],
    ]

    for reg, rev, mar, stat in data:
        row_cells = table.add_row().cells
        row_cells[0].text = reg
        row_cells[1].text = rev
        row_cells[2].text = mar
        row_cells[3].text = stat

    # Section 3
    doc_docx.add_heading("3. Internal Control Recommendations", level=1)
    doc_docx.add_paragraph(
        "Audit recommendation: Institute automated daily ERP margin checks to prevent unapproved price degradation. "
        "Reconcile booking volumes with billing ledger on a continuous 48-hour cycle to avoid end-of-quarter adjustments."
    )

    doc_docx.save(str(docx_path))
    print(f"Generated {docx_path.name}")


if __name__ == "__main__":
    generate_document_samples()
